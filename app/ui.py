from pathlib import Path
from io import BytesIO
import streamlit as st
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from app.database import (
    initialize_database,
    create_request,
    update_request_result,
    mark_request_failed,
    add_audit_log,
)
from app.rag.ingest import ingest_patient_documents
from app.workflow import build_workflow

# CONFIGURATION
st.set_page_config(
    page_title="PriorAuth AI",
    page_icon="⚕️",
    layout="wide",
    initial_sidebar_state="collapsed",
)
initialize_database()

# CSS
st.markdown(
    """
    <style>
    .block-container {
        max-width: 1250px;
        padding: 1.5rem 2rem 3rem;
    }
    div[data-testid="stVerticalBlock"] {
        gap: 0.55rem;
    }
    div[data-testid="stFileUploader"] section {
        padding: 0.6rem;
        border-radius: 8px;
    }
    button[kind="primary"] {
        height: 2.8rem;
        border-radius: 8px;
        font-weight: 700;
    }
    h1 {
        margin-top: 0 !important;
        margin-bottom: 0.1rem !important;
    }
    h2, h3 {
        margin-top: 0.7rem !important;
        margin-bottom: 0.6rem !important;
    }
    hr {
        margin: 1rem 0 !important;
    }

    </style>
    """,
    unsafe_allow_html=True,
)

# HEADER
st.title("⚕️ PriorAuth AI")
st.caption(
    "Evidence-grounded prior authorization assistant"
)
st.caption(
    "Guideline Analysis  •  Patient Evidence Retrieval  •  "
    "Evidence Matching  •  Validation"
)

# DOCUMENT GENERATOR
def set_cell_shading(cell, fill):
    """
    Add background shading to a Word table cell.
    """
    tc_pr = cell._tc.get_or_add_tcPr()

    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)

    tc_pr.append(shd)

def set_cell_text_bold(cell):
    """
    Make all text inside a Word cell bold.
    """
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = True

def build_pa_docx(
    request_id,
    patient_id,
    treatment,
    insurer,
    guideline,
    matches,
    draft,
    final_decision,
):
    from datetime import datetime

    """
    Generate an editable Microsoft Word PA draft.

    The document is intentionally a draft / decision-support
    artifact and does not represent payer approval.
    """

    document = Document()

    # PAGE MARGINS
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    # DEFAULT FONT
    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)

    # TITLE
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = title.add_run("PRIOR AUTHORIZATION REQUEST")
    run.bold = True
    run.font.size = Pt(18)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = subtitle.add_run("AI-Assisted Draft — Human Review Required")
    run.italic = True
    run.font.size = Pt(10)

    # Generation Date & Time
    generated_on = document.add_paragraph()
    generated_on.alignment = WD_ALIGN_PARAGRAPH.CENTER

    timestamp = datetime.now().strftime("%d %B %Y, %I:%M %p")
    run = generated_on.add_run(f"Generated on: {timestamp}")
    run.font.size = Pt(9)
    
    document.add_paragraph()

    # REQUEST INFORMATION
    document.add_heading("1. Request Information", level=1)

    table = document.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    request_rows = [
        ("Request ID", str(request_id)),
        ("Patient ID", patient_id),
        ("Requested Treatment", treatment),
        ("Insurance Provider", insurer),
    ]
    for row, (label, value) in zip(table.rows, request_rows):
        row.cells[0].text = label
        row.cells[1].text = value

        set_cell_shading(row.cells[0], "EAF2F8")
        set_cell_text_bold(row.cells[0])

    document.add_paragraph()

    # FINAL DECISION
    document.add_heading("2. Requirement Assessment", level=1)

    decision_paragraph = document.add_paragraph()

    if final_decision == "PASS":
        decision_text = (
            "REQUIREMENTS SATISFIED — "
            "All mandatory policy requirements are supported "
            "by the available patient evidence."
        )
    elif final_decision == "FAIL":
        decision_text = (
            "REQUIREMENTS NOT SATISFIED — "
            "The available patient evidence affirmatively "
            "demonstrates that one or more mandatory requirements "
            "are not satisfied."
        )
    else:
        decision_text = (
            "HUMAN REVIEW REQUIRED — "
            "The available documentation is insufficient to "
            "determine whether all mandatory requirements are satisfied."
        )

    run = decision_paragraph.add_run(decision_text)
    run.bold = True
    run.font.size = Pt(11)

    document.add_paragraph(
        "This assessment is an evidence-based decision-support result. "
        "It does not represent an insurance payer approval or denial."
    )

    # REQUIREMENT TABLE
    requirement_table = document.add_table(
        rows=1,
        cols=4,
    )

    requirement_table.style = "Table Grid"
    requirement_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    headers = [
        "Requirement",
        "Status",
        "Mandatory",
        "Assessment",
    ]

    for i, header in enumerate(headers):
        requirement_table.rows[0].cells[i].text = header
        set_cell_shading(
            requirement_table.rows[0].cells[i],
            "D9EAF7",
        )
        set_cell_text_bold(
            requirement_table.rows[0].cells[i]
        )

    for requirement in guideline.requirements:

        match = next(
            (
                item
                for item in matches
                if item.requirement_id
                == requirement.requirement_id
            ),
            None,
        )

        if match is None:
            continue

        row = requirement_table.add_row()

        row.cells[0].text = (
            f"{requirement.requirement_id}\n"
            f"{requirement.description}"
        )

        row.cells[1].text = match.status

        row.cells[2].text = (
            "Yes"
            if requirement.mandatory
            else "No"
        )

        row.cells[3].text = match.reasoning

    document.add_paragraph()

    # PATIENT DIAGNOSIS
    document.add_heading("3. Diagnosis", level=1)

    document.add_paragraph(
        draft.diagnosis
    )

    # CLINICAL JUSTIFICATION
    document.add_heading(
        "4. Clinical Justification",
        level=1,
    )
    document.add_paragraph(
        draft.clinical_justification
    )

    # PREVIOUS TREATMENTS
    document.add_heading(
        "5. Previous Treatments",
        level=1,
    )
    if draft.previous_treatments:

        for treatment_item in draft.previous_treatments:
            document.add_paragraph(
                treatment_item,
                style="List Bullet",
            )
    else:
        document.add_paragraph(
            "No previous treatments were documented."
        )

    # MISSING INFORMATION

    document.add_heading(
        "6. Missing Information",
        level=1,
    )
    if draft.missing_information:
        for missing_item in draft.missing_information:

            document.add_paragraph(
                missing_item,
                style="List Bullet",
            )
    else:
        document.add_paragraph(
            "No missing information identified."
        )

    # EVIDENCE TRACEABILITY
    document.add_heading(
        "7. Evidence Traceability",
        level=1,
    )
    document.add_paragraph(
        "Patient-record sources supporting "
        "each requirement assessment."
    )
    for requirement in guideline.requirements:

        match = next(
            (
                item
                for item in matches
                if item.requirement_id
                == requirement.requirement_id
            ),
            None,
        )

        if match is None:
            continue

        paragraph = document.add_paragraph()
        run = paragraph.add_run(
            f"{requirement.requirement_id} — {match.status}"
        )

        run.bold = True

        for evidence in match.evidence:

            evidence_paragraph = document.add_paragraph(
                style="List Bullet",
            )
            page = (
                f"Page {evidence.page}"
                if evidence.page
                else "Page unavailable"
            )

            raw_reference = getattr(
                evidence,
                "reference",
                None,
            )

            # Avoid redundant references
            reference = ""

            if raw_reference:
                normalized_reference = str(
                    raw_reference
                ).strip().lower()

                page_reference = (
                    f"page {evidence.page}".lower()
                    if evidence.page
                    else ""
                )
                if normalized_reference != page_reference:
                    reference = (
                        f" | {raw_reference}"
                    )
            evidence_paragraph.add_run(
                f"{evidence.document} | "
                f"{page}"
                f"{reference}\n"
                f"{evidence.text}"
            )
    # --------------------------------------------------------
    # REVIEWER SECTION
    # --------------------------------------------------------

    document.add_heading(
        "8. Reviewer Notes",
        level=1,
    )

    document.add_paragraph(
        "Reviewer comments:"
    )

    document.add_paragraph(
        "\n\n\n"
    )

    document.add_paragraph(
        "Reviewer Name: ______________________________"
    )

    document.add_paragraph(
        "Date: ____________________"
    )

    document.add_paragraph(
        "Signature: _________________________________"
    )

    # --------------------------------------------------------
    # FOOTER
    # --------------------------------------------------------

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run(
        "PriorAuth AI — AI-assisted decision-support artifact. "
        "Human review required before submission."
    )
    footer_run.font.size = Pt(8)

    # --------------------------------------------------------
    # SAVE TO MEMORY
    # --------------------------------------------------------

    output = BytesIO()
    document.save(output)
    output.seek(0)
    return output

# REQUEST INPUT
st.markdown("### Prior Authorization Request")
left, right = st.columns(2, gap="medium")

# LEFT
with left:

    with st.container(border=True):
        st.markdown("**Patient & Request**")

        patient_id = st.text_input(
            "Patient ID",
            placeholder="P001",
        )
        treatment = st.text_input(
            "Requested Treatment",
            placeholder="Drug X",
        )
        insurer = st.text_input(
            "Insurance Provider",
            placeholder="Demo Insurance",
        )

# RIGHT
with right:

    with st.container(border=True):

        st.markdown("**Documents**")

        guideline_file = st.file_uploader(
            "Insurance guideline",
            type=["pdf"],
            help="Upload the insurer's clinical policy.",
        )
        patient_files = st.file_uploader(
            "Patient medical records",
            type=["pdf"],
            accept_multiple_files=True,
            help="Upload one or more patient records.",
        )

# ============================================================
# GENERATE BUTTON
# ============================================================

run = st.button(
    "⚡ Generate Prior Authorization",
    type="primary",
    use_container_width=True,
)

# ============================================================
# WORKFLOW
# ============================================================

if run:

    # INPUT VALIDATION
    if not patient_id.strip():
        st.error("Enter a Patient ID.")
        st.stop()
    if not treatment.strip():
        st.error("Enter the requested treatment.")
        st.stop()
    if not insurer.strip():
        st.error("Enter the insurance provider.")
        st.stop()
    if guideline_file is None:
        st.error(
            "Upload the insurance guideline PDF."
        )
        st.stop()
    if not patient_files:
        st.error(
            "Upload at least one patient medical record."
        )
        st.stop()

    # CREATE REQUEST
    request_id = create_request(
        patient_id=patient_id,
        treatment=treatment,
        insurer=insurer,
    )
    add_audit_log(
        request_id,
        "DOCUMENTS_RECEIVED",
        f"{len(patient_files)} patient document(s) uploaded.",
    )

    # SAVE GUIDELINE
    guideline_dir = Path("data/guidelines")
    guideline_dir.mkdir(
        parents=True,
        exist_ok=True,
    )
    guideline_path = (
        guideline_dir
        / guideline_file.name
    )
    with open(
        guideline_path,
        "wb",
    ) as file:
        file.write(
            guideline_file.getbuffer()
        )

    # SAVE PATIENT DOCUMENTS
    patient_dir = (
        Path("data/patients")
        / patient_id
    )
    patient_dir.mkdir(
        parents=True,
        exist_ok=True,
    )
    current_patient_paths = []
    for uploaded_file in patient_files:
        file_path = (
            patient_dir
            / uploaded_file.name
        )
        with open(
            file_path,
            "wb",
        ) as file:
            file.write(
                uploaded_file.getbuffer()
            )
        current_patient_paths.append(
            str(file_path)
        )

    # RAG INGESTION
    with st.spinner(
        "Indexing patient records..."
    ):
        ingest_patient_documents(
            patient_id=patient_id,
            request_id=request_id,
            pdf_paths=current_patient_paths,
        )
    add_audit_log(
        request_id,
        "RAG_INGESTION_COMPLETED",
        "Patient records indexed successfully.",
    )

    # RUN WORKFLOW
    try:
        workflow = build_workflow()
        initial_state = {
            "request_id": request_id,
            "patient_id": patient_id,
            "treatment": treatment,
            "insurer": insurer,
            "guideline_path": str(
                guideline_path
            ),
        }
        with st.spinner(
            "Running PriorAuth AI..."
        ):
            result = workflow.invoke(
                initial_state
            )

        # EXTRACT RESULTS
        draft = result["draft"]
        matches = result["evidence_matches"]
        guideline = result["guideline"]
        raw_validation_result = result.get(
            "validation_result",
            "",
        )

        # REQUIREMENT COUNTS
        total_requirements = len(matches)
        satisfied = sum(
            1
            for match in matches
            if match.status == "SATISFIED"
        )
        not_satisfied = sum(
            1
            for match in matches
            if match.status == "NOT_SATISFIED"
        )
        unknown = sum(
            1
            for match in matches
            if match.status == "UNKNOWN"
        )

        # MANDATORY REQUIREMENTS
        mandatory_ids = {
            req.requirement_id
            for req in guideline.requirements
            if req.mandatory
        }
        mandatory_not_satisfied = [
            match
            for match in matches
            if (
                match.requirement_id
                in mandatory_ids
                and match.status
                == "NOT_SATISFIED"
            )
        ]
        mandatory_unknown = [
            match
            for match in matches
            if (
                match.requirement_id
                in mandatory_ids
                and match.status
                == "UNKNOWN"
            )
        ]

        # DETERMINISTIC FINAL DECISION
        if total_requirements == 0:
            final_decision = "REVIEW"
        elif mandatory_not_satisfied:
            final_decision = "FAIL"
        elif mandatory_unknown:
            final_decision = "REVIEW"
        else:
            final_decision = "PASS"

        # SAVE RESULT TO DATABASE
        update_request_result(
            request_id=request_id,
            final_decision=final_decision,
            matched_requirements=satisfied,
            failed_requirements=not_satisfied,
            unknown_requirements=unknown,
        )
        add_audit_log(
            request_id,
            "WORKFLOW_COMPLETED",
            (
                f"Decision={final_decision}; "
                f"Matched={satisfied}; "
                f"NotSatisfied={not_satisfied}; "
                f"Unknown={unknown}"
            ),
        )

        # SUCCESS MESSAGE
        st.success(
            "Prior Authorization analysis completed."
        )

        # REVIEW HEADER
        st.markdown(
            f"## Review — Request #{request_id}"
        )
        st.caption(
            f"Patient: {patient_id}  •  "
            f"Treatment: {treatment}  •  "
            f"Insurer: {insurer}"
        )

        # SUMMARY METRICS
        c1, c2, c3, c4 = st.columns(4)
        c1.metric(
            "Requirements",
            total_requirements,
        )
        c2.metric(
            "Matched",
            satisfied,
        )
        c3.metric(
            "Not Met",
            not_satisfied,
        )
        c4.metric(
            "Unknown",
            unknown,
        )

        # SUMMARY BANNER
        if final_decision == "PASS":
            st.success(
                "🟢 REQUIREMENTS SATISFIED"
            )
            st.caption(
                "All mandatory policy requirements are "
                "supported by the available patient evidence."
            )
        elif final_decision == "FAIL":
            st.error(
                "🔴 REQUIREMENTS NOT SATISFIED"
            )
            st.caption(
                "The available patient evidence affirmatively "
                "demonstrates that one or more mandatory "
                "requirements are not satisfied."
            )
        else:
            st.warning(
                "🟡 HUMAN REVIEW REQUIRED — "
                "INSUFFICIENT DOCUMENTATION"
            )
            st.caption(
                "The system does not deny the request. "
                "The available evidence is insufficient to "
                "determine whether all mandatory requirements "
                "are satisfied. Additional documentation "
                "is required."
            )

        # REQUIREMENT ANALYSIS
        st.divider()
        st.markdown(
            "## Requirement Analysis"
        )
        for requirement in guideline.requirements:
            match = next(
                (
                    item
                    for item in matches
                    if item.requirement_id
                    == requirement.requirement_id
                ),
                None,
            )
            if match is None:
                continue

            # STATUS
            if match.status == "SATISFIED":
                status = "✅ MATCHED"
            elif match.status == "NOT_SATISFIED":
                status = "❌ NOT MATCHED"
            else:
                status = "⚠️ UNKNOWN"

            # CARD
            with st.container(
                border=True
            ):
                st.markdown(
                    f"### "
                    f"{requirement.requirement_id} "
                    f"— {status}"
                )
                st.write(
                    requirement.description
                )
                mandatory_text = (
                    "Mandatory"
                    if requirement.mandatory
                    else "Non-mandatory"
                )
                st.caption(
                    f"{mandatory_text} • "
                    f"{requirement.source_document} • "
                    f"Page "
                    f"{requirement.source_page or 'N/A'}"
                )
                if requirement.condition:
                    st.caption(
                        f"Condition: "
                        f"{requirement.condition}"
                    )
                else:
                    st.caption(
                        "Condition: None"
                    )

                # ASSESSMENT
                st.markdown(
                    f"**Assessment:** "
                    f"{match.reasoning}"
                )

                # UNKNOWN
                if match.status == "UNKNOWN":
                    st.warning(
                        "Additional evidence is required "
                        "to determine whether this requirement "
                        "is satisfied."
                    )

                # NOT SATISFIED
                elif (
                    match.status
                    == "NOT_SATISFIED"
                ):
                    st.error(
                        "The available evidence indicates "
                        "that this requirement is not satisfied."
                    )

                # EVIDENCE
                if match.evidence:
                    st.markdown(
                        "**Supporting Evidence**"
                    )
                    for evidence in match.evidence:
                        page = (
                            f"p.{evidence.page}"
                            if evidence.page
                            else "page unavailable"
                        )
                        reference = getattr(
                            evidence,
                            "reference",
                            None,
                        )
                        if reference:
                            source_line = (
                                f"📄 **{evidence.document}** "
                                f"({page}) • "
                                f"{reference}"
                            )
                        else:
                            source_line = (
                                f"📄 **{evidence.document}** "
                                f"({page})"
                            )
                        st.info(
                            f"{source_line}\n\n"
                            f"> {evidence.text}"
                        )
                else:
                    st.warning(
                        "No supporting patient evidence "
                        "was found."
                    )

        # FINAL DECISION
        st.divider()
        st.markdown(
            "## Final Decision"
        )
        if final_decision == "PASS":
            st.success(
                "🟢 REQUIREMENTS SATISFIED"
            )
            st.caption(
                "All mandatory requirements are supported "
                "by available patient evidence."
            )
        elif final_decision == "FAIL":
            st.error(
                "🔴 REQUIREMENTS NOT SATISFIED"
            )
            st.caption(
                "The available evidence demonstrates that "
                "one or more mandatory requirements are not met."
            )
            if mandatory_not_satisfied:
                st.markdown(
                    "**Mandatory requirements that are not satisfied:**"
                )
                for match in mandatory_not_satisfied:
                    requirement = next(
                        (
                            req
                            for req in guideline.requirements
                            if req.requirement_id
                            == match.requirement_id
                        ),
                        None,
                    )
                    if requirement:
                        st.error(
                            f"**{requirement.requirement_id}:** "
                            f"{requirement.description}"
                        )

        else:
            st.warning(
                "🟡 HUMAN REVIEW REQUIRED"
            )
            st.caption(
                "The available documentation is insufficient "
                "to determine whether all mandatory requirements "
                "are satisfied."
            )
            if mandatory_unknown:
                st.markdown(
                    "**Additional documentation needed:**"
                )
                for match in mandatory_unknown:
                    requirement = next(
                        (
                            req
                            for req in guideline.requirements
                            if req.requirement_id
                            == match.requirement_id
                        ),
                        None,
                    )
                    if requirement:
                        st.warning(
                            f"**{requirement.requirement_id}:** "
                            f"{requirement.description}"
                        )

        # GENERATED PA
        st.divider()
        st.markdown(
            "## Generated Prior Authorization"
        )
        with st.container(
            border=True
        ):
            st.markdown(
                f"### {draft.treatment}"
            )
            st.caption(
                f"Patient: {draft.patient_id} • "
                f"Insurer: {draft.insurer}"
            )
            st.markdown(
                "**Diagnosis**"
            )
            st.write(
                draft.diagnosis
            )
            st.markdown(
                "**Clinical Justification**"
            )
            st.write(
                draft.clinical_justification
            )
            st.markdown(
                "**Previous Treatments**"
            )
            if draft.previous_treatments:
                for item in draft.previous_treatments:

                    st.write(
                        f"• {item}"
                    )
            else:
                st.caption(
                    "No previous treatments documented."
                )
            if draft.missing_information:

                st.warning(
                    "Missing Information"
                )
                for item in draft.missing_information:

                    st.write(
                        f"• {item}"
                    )

        # DOWNLOADABLE EDITABLE PA
        st.markdown(
            "### 📄 Download Editable PA Draft"
        )
        st.caption(
            "The Word document is editable and can be reviewed, "
            "modified, and completed before submission."
        )
        docx_file = build_pa_docx(
            request_id=request_id,
            patient_id=patient_id,
            treatment=treatment,
            insurer=insurer,
            guideline=guideline,
            matches=matches,
            draft=draft,
            final_decision=final_decision,
        )
        filename = (
            f"PriorAuth_{patient_id}_"
            f"Request_{request_id}.docx"
        )
        st.download_button(
            label="⬇️ Download Editable Word PA Form",
            data=docx_file,
            file_name=filename,
            mime=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
            use_container_width=True,
        )

        # VALIDATION
        st.divider()
        st.markdown(
            "## Validation"
        )
        st.write(
            raw_validation_result
        )

        # AUDIT / REQUEST DETAILS
        with st.expander(
            "Audit & Request Details"
        ):
            st.caption(
                f"Request ID: {request_id}"
            )
            st.caption(
                f"Patient ID: {patient_id}"
            )
            st.caption(
                f"Treatment: {treatment}"
            )
            st.caption(
                f"Insurer: {insurer}"
            )
            st.caption(
                f"Guideline: "
                f"{guideline_file.name}"
            )
            st.caption(
                f"Patient documents: "
                f"{len(patient_files)}"
            )
            st.caption(
                "Patient records: "
                + ", ".join(
                    file.name
                    for file in patient_files
                )
            )
            st.caption(
                f"Deterministic final decision: "
                f"{final_decision}"
            )
            st.caption(
                f"Matched: {satisfied} | "
                f"Not satisfied: {not_satisfied} | "
                f"Unknown: {unknown}"
            )

    # ERROR HANDLING
    except Exception as e:

        mark_request_failed(
            request_id=request_id,
            error_message=str(e),
        )
        add_audit_log(
            request_id,
            "WORKFLOW_FAILED",
            str(e),
        )
        st.error(
            "Prior Authorization workflow failed."
        )
        st.exception(e)